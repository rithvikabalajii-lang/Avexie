"""Local document artifact tools."""

from __future__ import annotations

import hashlib
import io
import re
from uuid import uuid4

from docx import Document

from avexie.models.files import FileForm, Files
from avexie.storage.provider import Storage
from avexie.utils.json_codec import JSONCodec


async def generate_word_document(
    title: str,
    body: str,
    filename: str = 'approval-note.docx',
    __user__: dict | None = None,
) -> str:
    """Create a downloadable Word document entirely on the local AVEXIE host.

    Use this after extracting or analysing a report when the user requests an
    approval note, report, memo, or other Word artifact.

    :param title: Document title
    :param body: Complete document text; blank lines create paragraphs
    :param filename: Output filename ending in .docx
    :return: JSON containing the local file id and authenticated download URL
    """
    if not __user__ or not __user__.get('id'):
        return JSONCodec.dumps({'error': 'User context not available'})

    safe_name = re.sub(r'[^A-Za-z0-9._ -]+', '-', filename).strip(' .-') or 'approval-note.docx'
    if not safe_name.lower().endswith('.docx'):
        safe_name += '.docx'

    document = Document()
    document.add_heading(title.strip() or 'Document', level=0)
    for paragraph in re.split(r'\n\s*\n', body.strip()):
        document.add_paragraph(paragraph.strip())

    buffer = io.BytesIO()
    document.save(buffer)
    contents = buffer.getvalue()
    file_id = str(uuid4())
    storage_name = f'{file_id}_{safe_name}'
    _, path = Storage.upload_file(io.BytesIO(contents), storage_name, {'Avexie-File-Id': file_id})
    await Files.insert_new_file(
        __user__['id'],
        FileForm(
            id=file_id,
            hash=hashlib.sha256(contents).hexdigest(),
            filename=safe_name,
            path=path,
            data={'generated_by': 'generate_word_document'},
            meta={
                'name': safe_name,
                'content_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                'size': len(contents),
                'local_only': True,
            },
        ),
    )
    return JSONCodec.dumps(
        {
            'status': 'created',
            'file_id': file_id,
            'filename': safe_name,
            'download_url': f'/api/v1/files/{file_id}/content?attachment=true',
            'external_calls': 0,
        }
    )
